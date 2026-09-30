"""PASO 1 - LECTURA de archivos de datos: CSV, TXT, JSON y Excel (.xlsx), o desde una URL/API.

Sin importar el formato, siempre devuelve lo mismo: una lista de FilaLeida.
Así el resto del sistema no necesita saber de qué tipo de archivo vinieron los datos.
Esta parte es genérica: normalmente NO hace falta cambiarla al adaptar la plantilla.
"""

from __future__ import annotations

import csv
import io
import json
import urllib.error
import urllib.parse
import urllib.request
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from app.utils.conversiones import es_vacio, normalizar_clave

FORMATOS_SOPORTADOS = ("csv", "txt", "json", "xlsx")
SEPARADORES_POSIBLES = ";,\t|"
VALORES_DE_MAS = "_valores_de_mas"  # marca de filas con más valores que columnas


class ErrorLectura(Exception):
    """El archivo completo no se puede procesar (vacío, formato inválido, faltan columnas...)."""


@dataclass
class FilaLeida:
    numero: int  # número de fila/elemento en el archivo original (para los mensajes de error)
    datos: dict[str, Any]  # nombre_de_columna_normalizado -> valor tal como vino


def extension_de(nombre_archivo: str) -> str:
    return Path(nombre_archivo).suffix.lower().lstrip(".")


def leer_archivo(nombre_archivo: str, contenido: bytes) -> list[FilaLeida]:
    extension = extension_de(nombre_archivo)
    if extension not in FORMATOS_SOPORTADOS:
        formatos = ", ".join(f".{f}" for f in FORMATOS_SOPORTADOS)
        raise ErrorLectura(f"El formato '.{extension}' no está soportado. Usá {formatos}.")
    if not contenido or not contenido.strip():
        raise ErrorLectura("El archivo está vacío.")

    if extension in ("csv", "txt"):
        filas = leer_csv(decodificar(contenido))
    elif extension == "json":
        filas = leer_json(decodificar(contenido))
    else:
        filas = leer_xlsx(contenido)

    if not filas:
        raise ErrorLectura("El archivo no tiene filas con datos (solo encabezados).")
    return filas


def descargar(url: str, limite_bytes: int) -> tuple[str, bytes]:
    """Descarga datos de una dirección web o API (http/https). Devuelve (nombre, contenido).

    Si la dirección no termina en .csv/.json/..., el formato se deduce del tipo de contenido.
    """
    url = url.strip()
    if not url.lower().startswith(("http://", "https://")):
        raise ErrorLectura("La dirección tiene que empezar con http:// o https://.")
    try:
        pedido = urllib.request.Request(url, headers={"User-Agent": "sistema-oniet/1.0"})
        with urllib.request.urlopen(pedido, timeout=15) as respuesta:
            contenido = respuesta.read(limite_bytes + 1)
            tipo = respuesta.headers.get_content_type()
    except urllib.error.HTTPError as error:  # el servidor respondió con un error (404, 500...)
        raise ErrorLectura(
            f"La dirección respondió con un error ({error.code}). Revisá que esté bien escrita."
        ) from None
    except (OSError, ValueError):  # sin conexión, dirección inválida, tiempo agotado...
        raise ErrorLectura(
            "No se pudo descargar la dirección. Revisá que esté bien escrita y que haya conexión."
        ) from None
    if len(contenido) > limite_bytes:
        raise ErrorLectura("Lo descargado es demasiado grande.")

    nombre = Path(urllib.parse.urlparse(url).path).name or "datos"
    if extension_de(nombre) not in FORMATOS_SOPORTADOS:
        parece_json = "json" in tipo or contenido.lstrip()[:1] in (b"[", b"{")
        nombre = f"{nombre}.{'json' if parece_json else 'csv'}"
    return nombre, contenido


def decodificar(contenido: bytes) -> str:
    """Pasa de bytes a texto. Excel en Windows suele guardar los CSV en cp1252, no en UTF-8,
    y la opción "Texto Unicode" de Excel guarda en UTF-16."""
    if contenido.startswith((b"\xff\xfe", b"\xfe\xff")):
        return contenido.decode("utf-16")
    for codificacion in ("utf-8-sig", "cp1252"):
        try:
            return contenido.decode(codificacion)
        except UnicodeDecodeError:
            continue
    return contenido.decode("latin-1")  # latin-1 nunca falla


def detectar_separador(texto: str) -> str:
    """Elige el separador (';' ',' tab o '|') que más aparece en una línea.

    Se miran las primeras 10 líneas con contenido: así un título como "Ventas, marzo 2026"
    arriba del encabezado no confunde la detección.
    """
    lineas = [linea for linea in texto.splitlines() if linea.strip()][:10]
    mejor, mayor = ",", 0
    for separador in SEPARADORES_POSIBLES:
        cantidad = max((linea.count(separador) for linea in lineas), default=0)
        if cantidad > mayor:
            mejor, mayor = separador, cantidad
    return mejor


def leer_csv(texto: str) -> list[FilaLeida]:
    lector = csv.reader(io.StringIO(texto), delimiter=detectar_separador(texto))
    try:
        filas_crudas = [(lector.line_num, valores) for valores in lector]
    except csv.Error:
        raise ErrorLectura(
            f"El archivo está mal formado cerca de la línea {lector.line_num} "
            "(¿hay comillas sin cerrar?)."
        ) from None
    return _filas_con_encabezado(filas_crudas)


def leer_json(texto: str) -> list[FilaLeida]:
    """Acepta los formatos más comunes de archivos y APIs:

    - una lista de objetos:                    [{...}, {...}]
    - un objeto con una lista (aunque esté anidada): {"ventas": [...]}, {"data": {"items": [...]}}
    - GeoJSON (datos abiertos con mapas):      se toman las "properties" de cada "feature"
    """
    try:
        datos = json.loads(texto)
    except json.JSONDecodeError as error:
        raise ErrorLectura(f"El JSON no es válido (línea {error.lineno}): {error.msg}.") from None

    if isinstance(datos, dict) and datos.get("type") == "FeatureCollection":
        datos = [
            elemento.get("properties") or {}
            for elemento in datos.get("features", [])
            if isinstance(elemento, dict)
        ]
    elif isinstance(datos, dict):
        listas = _listas_de_objetos(datos)
        if not listas and any(valor == [] for valor in datos.values()):
            listas = [[]]  # {"ventas": []}: archivo válido pero sin filas
        if len(listas) != 1:
            raise ErrorLectura(
                "El JSON debe ser una lista de objetos o tener una sola lista de objetos adentro."
            )
        datos = listas[0]
    if not isinstance(datos, list):
        raise ErrorLectura("El JSON debe ser una lista de objetos.")

    filas = []
    for posicion, elemento in enumerate(datos, start=1):
        if not isinstance(elemento, dict):
            raise ErrorLectura(f"El elemento {posicion} del JSON no es un objeto {{...}}.")
        claves = _preparar_encabezados(list(elemento.keys()))  # rechaza "Precio" y "precio" juntos
        datos_fila = dict(zip(claves, elemento.values(), strict=True))
        filas.append(FilaLeida(numero=posicion, datos=datos_fila))
    return filas


def _listas_de_objetos(objeto: dict, profundidad: int = 3) -> list[list]:
    """Busca listas de objetos dentro de un JSON (hasta 3 niveles de profundidad)."""
    encontradas = []
    for valor in objeto.values():
        if isinstance(valor, list) and valor and all(isinstance(v, dict) for v in valor):
            encontradas.append(valor)
        elif isinstance(valor, dict) and profundidad > 1:
            encontradas.extend(_listas_de_objetos(valor, profundidad - 1))
    return encontradas


def leer_xlsx(contenido: bytes) -> list[FilaLeida]:
    """Lee la primera hoja del Excel (el encabezado se busca igual que en CSV)."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ErrorLectura(
            "Para leer Excel hay que instalar openpyxl: pip install openpyxl"
        ) from None
    try:
        libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception:  # openpyxl lanza errores distintos según qué esté dañado
        raise ErrorLectura("No se pudo abrir el Excel. ¿Es un archivo .xlsx válido?") from None

    try:
        filas_crudas = [
            (numero, list(valores))
            for numero, valores in enumerate(libro.active.iter_rows(values_only=True), start=1)
        ]
    finally:
        libro.close()
    return _filas_con_encabezado(filas_crudas)


def _filas_con_encabezado(filas_crudas: list[tuple[int, list]]) -> list[FilaLeida]:
    """Busca el encabezado y arma las filas de datos.

    El encabezado es la primera fila con 2 o más valores: así se saltean los títulos que suelen
    tener los archivos exportados ("Reporte de ventas marzo"). Si ninguna fila tiene 2 valores
    (archivo de una sola columna), se usa la primera fila con contenido.
    """
    con_contenido = [
        (numero, valores)
        for numero, valores in filas_crudas
        if not all(es_vacio(valor) for valor in valores)  # se ignoran las filas vacías
    ]
    if not con_contenido:
        return []
    posicion_encabezado = next(
        (
            posicion
            for posicion, (_, valores) in enumerate(con_contenido)
            if sum(not es_vacio(valor) for valor in valores) >= 2
        ),
        0,
    )
    encabezados = _preparar_encabezados(con_contenido[posicion_encabezado][1])
    return [
        FilaLeida(numero=numero, datos=_combinar(encabezados, valores))
        for numero, valores in con_contenido[posicion_encabezado + 1 :]
    ]


def _preparar_encabezados(valores) -> list[str]:
    encabezados = [normalizar_clave(valor) for valor in valores]
    repetidos = sorted({e for e in encabezados if e and encabezados.count(e) > 1})
    if repetidos:
        raise ErrorLectura(f"Hay columnas repetidas en el encabezado: {', '.join(repetidos)}.")
    return encabezados


def _combinar(encabezados: list[str], valores: list) -> dict[str, Any]:
    """Une encabezados y valores. Si a la fila le faltan valores, quedan como ''.

    Si le sobran valores con contenido, se marca con la clave VALORES_DE_MAS para que la
    validación lo informe (suele ser una coma decimal sin comillas en un CSV separado por comas).
    """
    datos = {
        clave: (valores[posicion] if posicion < len(valores) else "")
        for posicion, clave in enumerate(encabezados)
        if clave  # se ignoran columnas sin nombre
    }
    sobrantes = [valor for valor in valores[len(encabezados) :] if not es_vacio(valor)]
    if sobrantes:
        datos[VALORES_DE_MAS] = len(sobrantes)
    return datos
